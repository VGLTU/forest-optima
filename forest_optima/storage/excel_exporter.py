# -*- coding: utf-8 -*-
"""
Модуль генерации научно-технических отчетов в формате Microsoft Excel (.xlsx).
Использует библиотеку openpyxl со специализированным типографическим оформлением.
"""

from __future__ import annotations
from pathlib import Path
from typing import Sequence

from forest_optima.domain.models import DistrictEvaluation, EnterpriseEvaluation

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def export_to_excel(
    districts_res: Sequence[DistrictEvaluation],
    enterprises_res: Sequence[EnterpriseEvaluation],
    output_path: str | Path,
) -> Path:
    """
    Генерирует двухстраничную аналитическую книгу Excel с корпоративным стилем оформления:
    1. 'Рейтинг_Лесничеств' - многокритериальное ранжирование и смещение позиций.
    2. 'Оценка_Арендаторов' - матрица Эйлера-Венна и интегральный уровень устойчивости.

    Raises:
        ImportError: Если библиотека openpyxl не установлена в среде исполнения.
    """
    if not OPENPYXL_AVAILABLE:
        raise ImportError(
            "Пакет openpyxl не обнаружен. Установите зависимость: pip install openpyxl"
        )

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()

    # Палитра оформления
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True, color="000000")
    font_regular = Font(name="Calibri", size=11, color="000000")

    fill_header = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    fill_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_highlight = PatternFill(start_color="E8F8F5", end_color="E8F8F5", fill_type="solid")

    border_thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # =========================================================================
    # Лист 1: Рейтинг_Лесничеств
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Рейтинг_Лесничеств"
    ws1.views.sheetView[0].showGridLines = True
    ws1.freeze_panes = "A2"
    ws1.row_dimensions[1].height = 28

    headers1 = [
        "Наименование лесничества",
        "Запас древесины (тыс. м³)",
        "Разрешенная площадь (га)",
        "Лесные дороги (км)",
        "Базовая модель (EV_raw)",
        "Ранг в базовой модели",
        "Индекс Лес-Оптима (WSM)",
        "Индекс TOPSIS",
        "Скорр. ранг",
        "Смещение ранга",
    ]

    ws1.append(headers1)
    for col_idx in range(1, len(headers1) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = border_thin

    for row_idx, r in enumerate(districts_res, start=2):
        ws1.row_dimensions[row_idx].height = 21
        delta_str = f"+{r.rank_shift}" if r.rank_shift > 0 else str(r.rank_shift)

        row_data = [
            r.district.name,
            r.district.timber_stock_k_m3,
            r.district.allowed_area_ha,
            r.district.road_length_km,
            r.raw_score,
            r.raw_rank,
            r.wsm_score,
            r.topsis_score,
            r.optimal_rank,
            delta_str,
        ]
        ws1.append(row_data)

        row_fill = fill_even if row_idx % 2 == 0 else fill_odd
        for col_idx in range(1, len(row_data) + 1):
            cell = ws1.cell(row=row_idx, column=col_idx)
            cell.font = font_regular
            cell.fill = row_fill
            cell.border = border_thin

            if col_idx == 1:
                cell.alignment = align_left
            elif col_idx in [2, 3, 4, 5]:
                cell.alignment = align_right
                cell.number_format = "#,##0.0"
            elif col_idx in [7, 8]:
                cell.alignment = align_center
                cell.number_format = "0.0000"
                if col_idx == 7:
                    cell.fill = fill_highlight
            else:
                cell.alignment = align_center

            if col_idx == 9:
                cell.font = font_bold

    # Автоподбор ширины колонок с запасом безопасности
    for col in ws1.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 16)
    ws1.column_dimensions["A"].width = 24
    ws1.column_dimensions["G"].width = 26
    ws1.column_dimensions["H"].width = 18

    # =========================================================================
    # Лист 2: Оценка_Арендаторов
    # =========================================================================
    ws2 = wb.create_sheet(title="Оценка_Арендаторов")
    ws2.views.sheetView[0].showGridLines = True
    ws2.freeze_panes = "A2"
    ws2.row_dimensions[1].height = 28

    headers2 = [
        "Наименование предприятия",
        "Экономика (У_экон)",
        "Экология (У_экол)",
        "Социум (У_соц)",
        "Интеграл (0-15)",
        "Качественный уровень",
        "Зона матрицы Эйлера-Венна",
    ]

    ws2.append(headers2)
    for col_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_header
        cell.border = border_thin

    for row_idx, e in enumerate(enterprises_res, start=2):
        ws2.row_dimensions[row_idx].height = 22
        row_data = [
            e.enterprise.name,
            e.enterprise.economic_score,
            e.enterprise.ecological_score,
            e.enterprise.social_score,
            e.total_score,
            e.sustainability_level,
            e.euler_venn_zone,
        ]
        ws2.append(row_data)

        row_fill = fill_even if row_idx % 2 == 0 else fill_odd
        for col_idx in range(1, len(row_data) + 1):
            cell = ws2.cell(row=row_idx, column=col_idx)
            cell.font = font_regular
            cell.fill = row_fill
            cell.border = border_thin

            if col_idx in [1, 6, 7]:
                cell.alignment = align_left
            else:
                cell.alignment = align_center

            if col_idx in [2, 3, 4]:
                cell.number_format = "0.0"

            if col_idx == 5:
                cell.font = font_bold
                cell.fill = fill_highlight
                cell.number_format = "0.0"

    ws2.column_dimensions["A"].width = 34
    ws2.column_dimensions["B"].width = 20
    ws2.column_dimensions["C"].width = 20
    ws2.column_dimensions["D"].width = 20
    ws2.column_dimensions["E"].width = 20
    ws2.column_dimensions["F"].width = 38
    ws2.column_dimensions["G"].width = 50

    wb.save(path)
    return path
